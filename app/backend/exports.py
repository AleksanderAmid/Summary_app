"""Local exports generated from text, with no uploaded file metadata/codebook."""
from io import BytesIO
import re
import zipfile
from xml.sax.saxutils import escape

FORMATS = {
    "txt": ("text/plain; charset=utf-8", ".txt"),
    "md": ("text/markdown; charset=utf-8", ".md"),
    "pdf": ("application/pdf", ".pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
}


def available_formats():
    import importlib.util
    return {"txt": True, "md": True,
            "pdf": importlib.util.find_spec("reportlab") is not None,
            "docx": importlib.util.find_spec("docx") is not None}


def document_parts(record, kind):
    if kind == "summary":
        body = record.get("summary", "")
        if not body:
            raise ValueError("No summary is available.")
        if record.get("uncertainties"):
            body += "\n\nUnresolved questions\n" + "\n".join(record["uncertainties"])
        if record.get("transcription", {}).get("warnings"):
            body += "\n\nTranscription review notes\n" + "\n".join(record["transcription"]["warnings"])
        return "Clinical summary", body
    if kind == "pseudonymised":
        if not record.get("pseudonymisation", {}).get("implemented"):
            raise ValueError("This older run has no pseudonymised document. Process its source again.")
        units = record.get("evidence", [])
        body = "\n\n".join("[" + u["id"] + "] " + u.get("document", "") + "\n" + u["text"] for u in units)
        body = body or record["pseudonymised_text"]
        if record.get("transcription", {}).get("warnings"):
            body += "\n\nTranscription review notes\n" + "\n".join(record["transcription"]["warnings"])
        return "Pseudonymised documents", body
    raise ValueError("Choose summary, pseudonymised documents, or both.")


def _docx(title, body):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    style = document.styles["Title"]
    style.font.name, style.font.size = "Calibri", Pt(20)
    style.font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph(title, "Title")
    for block in re.split(r"\n\s*\n", body.strip()):
        document.add_paragraph(block)
    props = document.core_properties
    props.author = props.last_modified_by = ""
    props.title, props.subject, props.comments = title, "", ""
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(title, body):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path
    import reportlab
    # ReportLab ships Vera; embedding it supports Swedish and avoids system-font dependencies.
    name = "SmartDocVera"
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(Path(reportlab.__file__).parent / "fonts/Vera.ttf")))
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=61, leftMargin=61,
                            topMargin=58, bottomMargin=58, title=title, author="")
    heading = ParagraphStyle("Title", fontName=name, fontSize=19, leading=25, textColor=colors.black, spaceAfter=18)
    normal = ParagraphStyle("Body", fontName=name, fontSize=10.5, leading=15,
                            alignment=TA_LEFT, spaceAfter=9, splitLongWords=True)
    story = [Paragraph(escape(title), heading)]
    for block in re.split(r"\n\s*\n", body.strip()):
        story.append(Paragraph(escape(block).replace("\n", "<br/>"), normal))
    doc.build(story)
    return buffer.getvalue()


def render_document(title, body, format):
    # Remove XML-disallowed controls while preserving tabs/newlines and Unicode text.
    body = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", body)
    if format == "txt":
        return (title + "\n\n" + body + "\n").encode("utf-8")
    if format == "md":
        return ("# " + title + "\n\n" + body + "\n").encode("utf-8")
    if format == "docx":
        return _docx(title, body)
    if format == "pdf":
        return _pdf(title, body)
    raise ValueError("Unsupported export format.")


def export_record(record, kind, format):
    if format not in FORMATS:
        raise ValueError("Choose Word, PDF, TXT, or Markdown.")
    if not available_formats()[format]:
        raise ValueError("This format needs an export package. Run setup.bat and restart the app.")
    kinds = ["summary", "pseudonymised"] if kind == "both" else [kind]
    output = []
    for item in kinds:
        title, body = document_parts(record, item)
        filename = ("summary" if item == "summary" else "pseudonymised-documents") + FORMATS[format][1]
        output.append((filename, render_document(title, body, format)))
    if len(output) == 1:
        filename, content = output[0]
        return filename, FORMATS[format][0], content
    bundle = BytesIO()
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as archive:
        for filename, content in output:
            archive.writestr(filename, content)
    return "smartdoc-documents.zip", "application/zip", bundle.getvalue()


def preview_pdf(record, kind, page=1):
    """Render locally for webviews without a built-in PDF reader."""
    import fitz
    if kind not in ("summary", "pseudonymised"):
        raise ValueError("Preview one document at a time.")
    title, body = document_parts(record, kind)
    with fitz.open(stream=render_document(title, body, "pdf"), filetype="pdf") as pdf:
        if not 1 <= page <= len(pdf):
            raise ValueError("This PDF page does not exist.")
        image = pdf[page - 1].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        return image.tobytes("png"), len(pdf)
