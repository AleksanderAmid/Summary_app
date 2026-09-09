"""Recognised icons must be retained; small text and unknown images still use OCR."""
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
import document_symbols
import transcription

WARNING = Path(__file__).parent / "fixtures" / "red-exclamation.png"
DESCRIPTION = "[Symbol: utropstecken i röd cirkel]"


def encode(image):
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


class SymbolTests(unittest.TestCase):
    def test_reviewed_warning_is_preserved_without_any_ocr_in_every_mode(self):
        for mode in transcription.MODES:
            with self.subTest(mode=mode), patch.object(transcription, "read_scan") as scan:
                result = transcription._read_word_image(WARNING.read_bytes(), mode, None)
            self.assertEqual(result["text"], DESCRIPTION)
            self.assertEqual(result["method"], "embedded_symbol")
            self.assertEqual(result["symbol_id"], "red_exclamation")
            self.assertEqual(result["attempts"], 0)
            scan.assert_not_called()

    def test_lossless_reencoding_and_white_compositing_keep_the_match(self):
        with Image.open(WARNING) as original:
            rgba = original.convert("RGBA")
            white = Image.new("RGBA", rgba.size, "white")
            white.alpha_composite(rgba)
            for image in (rgba, white.convert("RGB")):
                encoded = encode(image)
                self.assertNotEqual(encoded, WARNING.read_bytes())
                result = document_symbols.recognize_symbol(encoded)
                self.assertEqual(result["text"], DESCRIPTION)

    def test_small_medication_text_is_never_accepted_based_on_dimensions(self):
        image = Image.new("RGB", (60, 60), "white")
        ImageDraw.Draw(image).text((3, 10), "5 mg", fill="black")
        data = encode(image)
        self.assertIsNone(document_symbols.recognize_symbol(data))
        expected = {"text": "5 mg", "method": "tesseract", "warnings": []}
        with patch.object(transcription, "read_scan", return_value=expected) as scan:
            result = transcription._read_word_image(data, "balanced", None)
        self.assertEqual(result, expected)
        scan.assert_called_once_with(data, "balanced", None)

    def test_changed_pixels_do_not_pass_the_exact_match(self):
        with Image.open(WARNING) as original:
            changed = original.convert("RGBA")
        changed.putpixel((30, 30), (0, 0, 255, 255))
        data = encode(changed)
        self.assertIsNone(document_symbols.recognize_symbol(data))
        with patch.object(transcription, "read_scan", side_effect=RuntimeError("Unreadable image")) as scan:
            with self.assertRaisesRegex(RuntimeError, "Unreadable image"):
                transcription._read_word_image(data, "balanced", None)
        scan.assert_called_once()

    def test_multiframe_image_cannot_be_accepted_from_just_its_first_frame(self):
        with Image.open(WARNING) as original:
            first = original.convert("RGBA")
        second = Image.new("RGBA", first.size, "white")
        ImageDraw.Draw(second).text((2, 10), "5 mg", fill="black")
        buffer = io.BytesIO()
        first.save(buffer, format="TIFF", save_all=True, append_images=[second])
        self.assertIsNone(document_symbols.recognize_symbol(buffer.getvalue()))


class WordOrderingTests(unittest.TestCase):
    def test_empty_outer_table_cells_keep_their_column_positions(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "table.docx"
            document = Document()
            table = document.add_table(rows=2, cols=3)
            table.cell(0, 1).text = "Resultat"
            table.cell(0, 2).text = "Enhet"
            table.cell(1, 0).text = "Hb"
            table.cell(1, 1).text = "136"
            document.save(path)
            result = transcription.transcribe_docx(path)
        self.assertEqual(result["full_text"], "\tResultat\tEnhet\nHb\t136\t")

    def test_inline_symbols_stay_between_their_surrounding_text_and_table_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "synthetic.docx"
            document = Document()
            paragraph = document.add_paragraph("Before icon ")
            paragraph.add_run().add_picture(str(WARNING), width=Inches(.1))
            paragraph.add_run(" after icon.")
            table = document.add_table(rows=2, cols=3)
            paragraph = table.cell(0, 0).paragraphs[0]
            paragraph.add_run("Before table icon ")
            paragraph.add_run().add_picture(str(WARNING), width=Inches(.1))
            paragraph.add_run(" after table icon")
            table.cell(0, 2).text = "500 mg"
            table.cell(1, 0).text = "Hb"
            table.cell(1, 1).text = "136"
            table.cell(1, 2).text = "g/L"
            document.save(path)
            with patch.object(transcription, "read_scan") as scan:
                result = transcription.transcribe_docx(path)
            scan.assert_not_called()
        text = result["full_text"]
        self.assertIn("Before icon \n"+DESCRIPTION+"\n after icon.", text)
        self.assertIn("Before table icon \n"+DESCRIPTION+"\n after table icon\t\t500 mg", text)
        self.assertIn("Hb\t136\tg/L", text)
        self.assertEqual(text.count(DESCRIPTION), 2)
        self.assertEqual([page["page_num"] for page in result["pages"]], [1, 2])
        self.assertEqual(len(result["warnings"]), 1)
        self.assertIn("2 embedded symbols", result["warnings"][0])
        self.assertTrue(result["review_required"])

    def test_text_image_is_read_in_place_among_native_text_and_recognised_icons(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "mixed.docx"
            document = Document()
            paragraph = document.add_paragraph("Before ")
            paragraph.add_run().add_picture(str(WARNING), width=Inches(.1))
            paragraph.add_run(" middle ")
            text_image = Image.new("RGB", (60, 60), "white")
            ImageDraw.Draw(text_image).text((3, 10), "5 mg", fill="black")
            paragraph.add_run().add_picture(io.BytesIO(encode(text_image)), width=Inches(.1))
            paragraph.add_run(" after")
            document.save(path)
            with patch.object(transcription, "read_scan", return_value={
                    "text": "5 mg", "method": "tesseract", "warnings": []}) as scan:
                result = transcription.transcribe_docx(path)
            scan.assert_called_once()
            text = result["full_text"]
            self.assertIn(DESCRIPTION+"\n middle \n5 mg\n after", text)
            self.assertEqual([page["method"] for page in result["pages"]], ["embedded_symbol", "tesseract"])
            with patch.object(transcription, "read_scan", side_effect=RuntimeError("Unreadable image")):
                with self.assertRaisesRegex(RuntimeError, "Word image 2: Unreadable image"):
                    transcription.transcribe_docx(path)


if __name__ == "__main__":
    unittest.main()
