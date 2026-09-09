"""Concurrency and merge tests with synthetic pages and deliberately reordered work."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
import io
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/"backend"))
import anonymization as anon
import concurrency
import ocr_engines
import pipeline
import privacy
import transcription as tr
import fitz
from PIL import Image


class FiveAtOnce:
    def __init__(self):
        self.barrier = threading.Barrier(5, timeout=10)
        self.lock = threading.Lock()
        self.active = self.peak = 0
        self.finished = []

    def run(self, number):
        with self.lock:
            self.active += 1
            self.peak = max(self.peak, self.active)
        try:
            self.barrier.wait()
            time.sleep((5-number % 5)*.01)
        finally:
            with self.lock:
                self.active -= 1
                self.finished.append(number)
        return number


def image_bytes():
    buffer=io.BytesIO()
    Image.new("RGB",(200,200),"gray").save(buffer,format="PNG")
    return buffer.getvalue()


class ParallelTests(unittest.TestCase):
    def test_pdf_has_five_active_pages_and_retains_order(self):
        tracker=FiveAtOnce()
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"five.pdf"
            with fitz.open() as doc:
                for _ in range(5):
                    doc.new_page().insert_image(fitz.Rect(10,10,500,700),stream=image_bytes())
                doc.save(path)
            def scan(data,*args):
                number=tracker.run(int(data))
                return {"text":f"Page value {number}","method":"test","warnings":[]}
            with patch.object(tr,"_page_to_png",side_effect=lambda page:str(page.number).encode()), \
                 patch.object(tr,"read_scan",side_effect=scan):
                result=tr.transcribe_pdf(path)
        self.assertEqual(tracker.peak,5)
        self.assertNotEqual(tracker.finished,list(range(5)))
        self.assertEqual([p["page_num"] for p in result["pages"]],[1,2,3,4,5])
        positions=[result["full_text"].index(f"Page value {i}") for i in range(5)]
        self.assertEqual(positions,sorted(positions))

    def test_scan_limit_is_shared_across_nested_jobs_and_released_on_error(self):
        tracker=FiveAtOnce()
        def scan(data,*args):
            number=tracker.run(data)
            if number==0:
                raise RuntimeError("Synthetic failure")
            return number
        with patch.object(tr,"_read_scan",side_effect=scan), ThreadPoolExecutor(max_workers=10) as pool:
            futures=[pool.submit(tr.read_scan,i) for i in range(10)]
            with self.assertRaisesRegex(RuntimeError,"Synthetic failure"):
                futures[0].result(timeout=15)
            self.assertEqual([f.result(timeout=15) for f in futures[1:]],list(range(1,10)))
        self.assertEqual(tracker.peak,5)
        self.assertEqual(len(tracker.finished),10)

    def test_five_uploaded_images_are_concurrent_and_pdf_stays_on_caller_thread(self):
        tracker=FiveAtOnce();caller=threading.get_ident();pdf_threads=[]
        def read(path,*args):
            if path.suffix==".pdf":
                pdf_threads.append(threading.get_ident())
                return {"full_text":"PDF"}
            return {"full_text":str(tracker.run(int(path.stem)))}
        paths=[Path(f"{i}.png") for i in range(5)]+[Path("last.pdf")]
        with patch.object(tr,"transcribe_file",side_effect=read):
            results=list(tr.transcribe_files(paths))
        self.assertEqual(tracker.peak,5)
        self.assertEqual([i for i,_ in results],list(range(6)))
        self.assertEqual([r["full_text"] for _,r in results],["0","1","2","3","4","PDF"])
        self.assertEqual(pdf_threads,[caller])

    def test_five_word_images_retain_their_places_among_paragraphs(self):
        from docx import Document
        tracker=FiveAtOnce();counter=iter(range(5));lock=threading.Lock()
        def scan(*args):
            with lock: number=next(counter)
            tracker.run(number)
            return {"text":"Scanned note", "method":"test","warnings":[]}
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"images.docx"
            doc=Document()
            for i in range(5):
                doc.add_paragraph(f"Before {i}")
                doc.add_picture(io.BytesIO(image_bytes()))
                doc.add_paragraph(f"After {i}")
            doc.save(path)
            with patch.object(tr,"read_scan",side_effect=scan):
                result=tr.transcribe_docx(path)
        self.assertEqual(tracker.peak,5)
        for i in range(5):self.assertIn(f"Before {i}\n\nScanned note\n\nAfter {i}",result["full_text"])
        self.assertEqual([p["page_num"] for p in result["pages"]],[1,2,3,4,5])

    def test_page_chunks_keep_offsets_and_do_not_split_short_pages_together(self):
        text="\n\n".join(f"--- Document {i} ---\n--- Sida 1 ---\nErik Exempelsson. Note {i}." for i in range(1,6))
        chunks=list(anon._page_chunks(text))
        self.assertEqual(len(chunks),5)
        for offset,chunk in chunks:self.assertEqual(text[offset:offset+len(chunk)],chunk)
        self.assertEqual([chunk.count("Erik") for _,chunk in chunks],[1]*5)
        long="--- Sida 1 ---\n"+"Clinical text "*600
        chunks=list(anon._page_chunks(long))
        self.assertEqual("".join(chunk for _,chunk in chunks),long)
        self.assertTrue(all(len(chunk)<=3500 for _,chunk in chunks))

    def test_parallel_identifier_results_share_one_deterministic_mapping(self):
        tracker=FiveAtOnce()
        source="\n\n".join(f"--- Sida {i+1} ---\nErik Exempelsson tar 500 mg. Note {i}" for i in range(5))
        def detect(chunk,model):
            number=int(chunk.rstrip()[-1]);tracker.run(number)
            start=chunk.index("Erik Exempelsson")
            # A stronger category found late must apply consistently to every page.
            kind="PATIENT_MALE" if number==0 else "PERSON"
            return [(start,start+len("Erik Exempelsson"),kind)],number
        with patch.object(anon,"llm_spans",side_effect=detect):
            result=anon.run_anonymization_stage(source)
        self.assertEqual(tracker.peak,5)
        self.assertEqual(result["passages"],5)
        self.assertEqual(result["discarded_suggestions"],10)
        self.assertEqual(result["text"].count("[NAME_PATIENT_MALE_01]"),5)
        self.assertEqual(result["text"].count("500 mg"),5)
        self.assertEqual(len(result["mapping"]),1)

    def test_failed_identifier_passage_is_counted_after_all_checks_complete(self):
        tracker=FiveAtOnce()
        source="\n\n".join(f"--- Sida {i+1} ---\nDate 2026-08-01. Note {i}" for i in range(5))
        def detect(chunk,model):
            number=int(chunk.rstrip()[-1]);tracker.run(number)
            if number==2:raise RuntimeError("Synthetic model failure")
            return [],0
        with patch.object(anon,"llm_spans",side_effect=detect):
            result=anon.run_anonymization_stage(source)
        self.assertEqual(len(tracker.finished),5)
        self.assertEqual(result["fallback_passages"],1)
        self.assertEqual(result["text"].count("[DATE_01]"),5)
        self.assertTrue(any("1 of 5" in warning for warning in result["warnings"]))

    def test_paddle_uses_five_separate_predictors_and_returns_worker_after_failure(self):
        tracker=FiveAtOnce();used=set();lock=threading.Lock()
        class Worker:
            def predict(self,png,timeout):
                with lock:used.add(id(self))
                if isinstance(png,int):tracker.run(png)
                if png=="fail":raise RuntimeError("Synthetic failure")
                return png
            def close(self):pass
        with patch.object(ocr_engines,"PaddleWorker",Worker):pool=ocr_engines.PaddlePool()
        with ThreadPoolExecutor(max_workers=5) as threads:
            self.assertEqual(list(threads.map(pool.predict,range(5))),list(range(5)))
        self.assertEqual(len(used),5)
        with self.assertRaises(RuntimeError):pool.predict("fail")
        self.assertEqual(pool.available.qsize(),5)
        self.assertEqual(pool.predict("next"),"next")
        pool.close()

    def test_summary_waits_for_every_page_and_identifier_result(self):
        scans=FiveAtOnce();identifiers=FiveAtOnce();number=iter(range(5));lock=threading.Lock()
        def scan(*args):
            with lock:i=next(number)
            scans.run(i)
            return {"text":f"Erik Exempelsson. Metformin 500 mg. Note {i}","method":"test","warnings":[]}
        def detect(chunk,model):
            i=int(chunk.rstrip()[-1]);identifiers.run(i)
            start=chunk.index("Erik Exempelsson")
            return [(start,start+len("Erik Exempelsson"),"PATIENT_MALE")],0
        def summarize(text,*args):
            self.assertEqual(len(scans.finished),5)
            self.assertEqual(len(identifiers.finished),5)
            self.assertEqual(text.count("[NAME_PATIENT_MALE_01]"),5)
            return {"summary":"[NAME_PATIENT_MALE_01]: Metformin 500 mg.","telemetry":{}}
        with tempfile.TemporaryDirectory() as folder, ExitStack() as stack:
            stack.enter_context(patch.object(pipeline,"UPLOAD_DIR",Path(folder)))
            stack.enter_context(patch.object(privacy,"PRIVATE_DIR",Path(folder)/"private"))
            stack.enter_context(patch.object(pipeline,"_jobs",{"test":{"stages":pipeline._new_stages()}}))
            stack.enter_context(patch.object(tr,"_read_scan",side_effect=scan))
            stack.enter_context(patch.object(anon,"llm_spans",side_effect=detect))
            generated=stack.enter_context(patch.object(pipeline.summarization,"summarize",side_effect=summarize))
            stack.enter_context(patch.object(pipeline.history,"save_record"))
            result=pipeline._execute("test",{"files":[{"filename":f"{i}.png","content":image_bytes()} for i in range(5)]})
            self.assertIn("Erik Exempelsson",result["summary"])
            self.assertNotIn("Erik",result["pseudonymised_text"])
            self.assertFalse(list((Path(folder)/"private").glob("*.enc")))
            generated.assert_called_once()


if __name__ == "__main__":unittest.main()
