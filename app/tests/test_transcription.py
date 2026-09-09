"""Regression tests with synthetic pages; no patient files or model downloads."""
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/"backend"))
import transcription as tr
import ocr_engines
import uploads
from PIL import Image
import fitz


def png(blank=False):
    from PIL import ImageDraw
    image=Image.new("RGB",(600,400),"white")
    if not blank:
        ImageDraw.Draw(image).text((30,40),"Synthetic page 123",fill="black")
    buffer=io.BytesIO();image.save(buffer,format="PNG")
    return buffer.getvalue()


def candidate(engine="tesseract", text="Metformin 500 mg. Ingen feber.", confidence=96):
    line={"text":text,"box":[10,10,400,40],"confidence":confidence}
    return {"engine":engine,"text":text,"confidence":confidence,"lines":[line],"words":[line],"seconds":.1}


class ScanTests(unittest.TestCase):
    def setUp(self):
        self.available=patch.object(ocr_engines,"status",return_value={"tesseract":True,"paddleocr":True})
        self.available.start();self.addCleanup(self.available.stop)
        self.prepare=patch.object(tr,"prepare_image",return_value=(Image.new("L",(600,400),255),{"blank":False,"rotation":0,"deskew":0,"warnings":[]}))
        self.prepare.start();self.addCleanup(self.prepare.stop)

    def test_short_high_confidence_page_needs_only_one_reader(self):
        with patch.object(ocr_engines,"tesseract",return_value=candidate()) as tess, \
             patch.object(ocr_engines,"paddle") as paddle, patch.object(tr,"vision_transcribe") as vision:
            result=tr.read_scan(png())
        self.assertIn("500 mg",result["text"])
        self.assertEqual(result["attempts"],1)
        tess.assert_called_once();paddle.assert_not_called();vision.assert_not_called()

    def test_thorough_crosscheck_flags_changed_dose_and_negation(self):
        with patch.object(ocr_engines,"tesseract",return_value=candidate()), \
             patch.object(ocr_engines,"paddle",return_value=candidate("paddleocr","Metformin 50 mg. Feber.")):
            result=tr.read_scan(png(),mode="thorough")
        self.assertTrue(result["review_required"])
        self.assertTrue(any("disagree" in w for w in result["warnings"]))
        self.assertEqual(result["engines"],["tesseract","paddleocr"])

    def test_spelling_disagreement_is_visible_even_when_numbers_match(self):
        with patch.object(ocr_engines,"tesseract",return_value=candidate(text="Återbesök 2026-10-01")), \
             patch.object(ocr_engines,"paddle",return_value=candidate("paddleocr","Aterbesök 2026-10-01")):
            result=tr.read_scan(png(),mode="thorough")
        self.assertTrue(any("wording" in w for w in result["warnings"]))

    def test_second_reader_failure_keeps_valid_first_reader_with_warning(self):
        with patch.object(ocr_engines,"tesseract",return_value=candidate()), \
             patch.object(ocr_engines,"paddle",side_effect=RuntimeError("timeout")):
            result=tr.read_scan(png(),mode="thorough")
        self.assertEqual(result["method"],"tesseract")
        self.assertTrue(result["review_required"])

    def test_unreadable_ocr_does_not_produce_a_successful_page(self):
        with patch.object(ocr_engines,"tesseract",return_value=candidate(confidence=10)), \
             patch.object(ocr_engines,"paddle",return_value=candidate("paddleocr",confidence=12)):
            with self.assertRaisesRegex(RuntimeError,"reliably"):
                tr.read_scan(png())

    def test_blank_page_never_calls_an_engine(self):
        self.prepare.stop()
        with patch.object(ocr_engines,"tesseract") as tess, patch.object(tr,"vision_transcribe") as vision:
            result=tr.read_scan(png(blank=True))
        self.assertEqual(result["method"],"blank")
        tess.assert_not_called();vision.assert_not_called()


class VisionTests(unittest.TestCase):
    def test_complete_short_vision_output_is_not_retried(self):
        with patch.object(tr.ollama_client,"post_json",return_value={"response":"Hb 136 g/L.","done":True,"done_reason":"stop"}) as call:
            self.assertEqual(tr.transcribe_image_bytes(png()),"Hb 136 g/L.")
        call.assert_called_once()

    def test_truncation_and_repetition_fail_after_bounded_retries(self):
        for result in ({"response":"A partial transcription","done_reason":"length"},
                       {"response":"5"*40,"done_reason":"stop"}):
            with patch.object(tr.ollama_client,"post_json",return_value=result) as call:
                with self.assertRaises(RuntimeError): tr.transcribe_image_bytes(png())
            self.assertEqual(call.call_count,2)

    def test_repeated_clinical_lines_are_preserved(self):
        text="Hb 136 g/L\nHb 136 g/L\nHb 136 g/L"
        with patch.object(tr.ollama_client,"post_json",return_value={"response":text,"done_reason":"stop"}):
            self.assertEqual(tr.transcribe_image_bytes(png()),text)


class FileTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def pdf(self,name,build):
        doc=fitz.open();build(doc);path=self.root/name;doc.save(path);doc.close();return path

    def test_short_native_page_does_not_use_ocr(self):
        path=self.pdf("short.pdf",lambda d:d.new_page().insert_text((50,50),"Hb 136 g/L"))
        with patch.object(tr,"read_scan") as scan:
            result=tr.transcribe_pdf(path)
        self.assertIn("Hb 136 g/L",result["full_text"])
        self.assertEqual(result["pages"][0]["method"],"native")
        scan.assert_not_called()

    def test_native_header_does_not_hide_scanned_body(self):
        def build(doc):
            page=doc.new_page()
            page.insert_text((40,40),"Synthetic patient journal header with more than fifty characters")
            page.insert_image(fitz.Rect(30,100,565,800),stream=png())
        path=self.pdf("mixed.pdf",build)
        with patch.object(tr,"read_scan",return_value={"text":"Scanned body 500 mg","method":"tesseract","warnings":[]} ) as scan:
            result=tr.transcribe_pdf(path)
        self.assertIn("Scanned body",result["full_text"]);scan.assert_called_once()

    def test_duplicate_scans_are_read_once_and_page_order_is_retained(self):
        def build(doc):
            for _ in range(3): doc.new_page().insert_image(fitz.Rect(30,30,565,800),stream=png())
        path=self.pdf("duplicates.pdf",build)
        with patch.object(tr,"read_scan",return_value={"text":"Synthetic scan","method":"tesseract","warnings":[]} ) as scan:
            result=tr.transcribe_pdf(path)
        scan.assert_called_once()
        self.assertEqual([p["page_num"] for p in result["pages"]],[1,2,3])
        self.assertEqual([p["cached"] for p in result["pages"]],[False,True,True])
        self.assertEqual(result["full_text"].count("Synthetic scan"),3)

    def test_empty_failed_scan_is_not_hidden_by_page_markers(self):
        path=self.pdf("blank.pdf",lambda d:d.new_page())
        result=tr.transcribe_pdf(path)
        self.assertEqual(result["full_text"],"")
        self.assertEqual(result["pages"][0]["method"],"blank")
        path=self.pdf("fail.pdf",lambda d:d.new_page().insert_image(fitz.Rect(30,30,565,800),stream=png()))
        with patch.object(tr,"read_scan",side_effect=RuntimeError("Unreadable scan")):
            with self.assertRaisesRegex(RuntimeError,"Page 1"):
                tr.transcribe_pdf(path)

    def test_word_table_preserves_rows_and_embedded_image_text(self):
        from docx import Document
        from docx.shared import Inches
        doc=Document();doc.add_paragraph("Synthetic journal")
        table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text="Hb";table.cell(0,1).text="136 g/L"
        table.cell(1,0).text="Kalium";table.cell(1,1).text="4,2 mmol/L"
        doc.add_picture(io.BytesIO(png()),width=Inches(2))
        path=self.root/'table.docx';doc.save(path)
        with patch.object(tr,"read_scan",return_value={"text":"Embedded note","method":"tesseract","warnings":[]}):
            result=tr.transcribe_docx(path)
        self.assertIn("Hb\t136 g/L",result["full_text"])
        self.assertIn("Kalium\t4,2 mmol/L",result["full_text"])
        self.assertIn("Embedded note",result["full_text"])

    def test_unicode_text_and_invalid_json_are_explicit(self):
        path=self.root/'note.txt';path.write_bytes("Återbesök 4,2 mmol/L".encode('utf-16'))
        self.assertEqual(tr.transcribe_txt(path)["full_text"],"Återbesök 4,2 mmol/L")
        path=self.root/'note.json';path.write_text(json.dumps({"pages":[{"text":123}]}))
        with self.assertRaises(ValueError):tr.transcribe_study_json(path)

    def test_prose_columns_and_lab_rows_keep_their_reading_order(self):
        lines=[]
        for i in range(3):
            lines.extend([{"text":"Left "+chr(65+i),"box":[0,i*30,100,i*30+20]},
                          {"text":"Right "+chr(65+i),"box":[180,i*30,280,i*30+20]}])
        text=tr._line_layout(lines)
        self.assertLess(text.index("Left C"),text.index("Right A"))
        lines=[{"text":"Hb","box":[0,0,30,20]},{"text":"136 g/L","box":[180,0,240,20]}]
        self.assertEqual(tr._line_layout(lines),"Hb\t136 g/L")

    def test_white_text_on_dark_background_is_not_marked_blank(self):
        from PIL import ImageOps
        with Image.open(io.BytesIO(png())) as image:
            dark=ImageOps.invert(image.convert("RGB"))
            buffer=io.BytesIO();dark.save(buffer,format="PNG")
        with patch.object(tr,"_orientation",return_value=0), patch.object(tr,"_deskew",side_effect=lambda image:(image,0)):
            image,metadata=tr.prepare_image(buffer.getvalue())
        self.assertFalse(metadata["blank"])
        self.assertGreater(image.getpixel((0,0)),200)

    def test_empty_transcription_json_does_not_create_only_markers(self):
        path=self.root/'empty.json';path.write_text(json.dumps({"pages":[{"text":""}]}))
        self.assertEqual(tr.transcribe_study_json(path)["full_text"],"")

    def test_exports_keep_transcription_review_notes(self):
        import exports
        record={"summary":"Test summary", "pseudonymised_text":"Test source", "pseudonymisation":{"implemented":True},
                "transcription":{"warnings":["Page 1 needs review."]}}
        for kind in ("summary","pseudonymised"):
            _,_,content=exports.export_record(record,kind,"txt")
            self.assertIn(b"Page 1 needs review.",content)

    def test_three_column_lab_table_keeps_label_value_and_unit_together(self):
        lines=[]
        for i,row in enumerate((("Analys","Resultat","Enhet"),("Hb","136","g/L"),("Kalium","4,2","mmol/L"),("CRP","<5","mg/L"))):
            for j,text in enumerate(row):
                lines.append({"text":text,"box":[j*160,i*30,j*160+90,i*30+20]})
        text=tr._line_layout(lines)
        self.assertIn("Hb\t136\tg/L",text)
        self.assertIn("CRP\t<5\tmg/L",text)
        self.assertNotEqual(tr.critical_tokens("CRP <5 mg/L"),tr.critical_tokens("CRP >5 mg/L"))
        self.assertNotEqual(tr.critical_tokens("-4,2"),tr.critical_tokens("4,2"))

    def test_lab_cells_with_different_top_edges_stay_on_their_own_row(self):
        lines=[]
        expected=[]
        for i,row in enumerate((("Analys","Resultat","Enhet"),("Hb","136","g/L"),("Kalium","4,2","mmol/L"),("CRP","<5","mg/L"))):
            for j,text in enumerate(row):
                y=i*30+(4,0,2)[j]
                lines.append({"text":text,"box":[j*160,y,j*160+90,y+20]})
            expected.append("\t".join(row))
        self.assertEqual(tr._line_layout(list(reversed(lines))),"\n".join(expected))

    def test_reading_mode_is_validated_at_api_boundary(self):
        self.assertEqual(uploads.parse_payload({"text":"Test","transcription_mode":"thorough"})['transcription_mode'],'thorough')
        with self.assertRaises(uploads.UploadError):uploads.parse_payload({"text":"Test","transcription_mode":{}})


if __name__ == '__main__':unittest.main()